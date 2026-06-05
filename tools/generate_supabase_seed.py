from __future__ import annotations

from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Iterable

import openpyxl


WORKBOOK_PATH = Path("/Users/kunle/Downloads/Kano Health Finance Tracker.xlsx")
OUTPUT_DIR = Path("supabase/seeds")


def q(value) -> str:
    if value is None:
        return "null"
    text = str(value)
    return "'" + text.replace("'", "''") + "'"


def money(value) -> str:
    if value is None or value == "":
        return "0.00"
    dec = Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return format(dec, "f")


def slug(value: str) -> str:
    chars = []
    last_dash = False
    for char in value.lower():
        if char.isalnum():
            chars.append(char)
            last_dash = False
        elif not last_dash:
            chars.append("-")
            last_dash = True
    return "".join(chars).strip("-")


def col_values(ws, col_idx: int) -> list[str]:
    values = []
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row, col_idx).value
        if value is not None and str(value).strip():
            values.append(str(value).strip())
    return values


def values_block(rows: Iterable[tuple]) -> str:
    return ",\n".join("  (" + ", ".join(row) + ")" for row in rows)


def infer_mda_type(name: str) -> str:
    lowered = name.lower()
    if "ministry" in lowered:
        return "Ministry"
    if "board" in lowered:
        return "Board"
    if "hospital" in lowered:
        return "Hospital"
    if "agency" in lowered:
        return "Agency"
    if "college" in lowered:
        return "College"
    if "school" in lowered:
        return "School"
    if "trust fund" in lowered or "fund" in lowered:
        return "Fund"
    if "centre" in lowered or "center" in lowered:
        return "Centre"
    return "Other"


def reference_insert(table: str, values: list[str]) -> str:
    rows = [(q(slug(value)), q(value)) for value in values]
    return f"""insert into public.{table} (slug, name)
values
{values_block(rows)}
on conflict (slug) do update
set name = excluded.name,
    active = true,
    updated_at = now();
"""


def status_seed() -> str:
    rows = [
        ("pending", "Pending", "Submitted by an MDA user and awaiting review"),
        ("approved", "Approved", "Reviewed and accepted for reporting"),
        ("processed", "Processed", "Accepted and reconciled or posted in the finance process"),
        ("rejected", "Rejected", "Reviewed and rejected; excluded from official reporting"),
    ]
    return f"""insert into public.entry_statuses (slug, name, description)
values
{values_block((q(a), q(b), q(c)) for a, b, c in rows)}
on conflict (slug) do update
set name = excluded.name,
    description = excluded.description,
    active = true,
    updated_at = now();
"""


def mda_seed(lookup_ws) -> str:
    mdas = []
    for row in range(2, lookup_ws.max_row + 1):
        code = lookup_ws.cell(row, 1).value
        name = lookup_ws.cell(row, 2).value
        if code and name:
            mdas.append((str(code).strip(), str(name).strip(), infer_mda_type(str(name).strip())))

    mda_types = sorted({mda_type for _, _, mda_type in mdas})
    mda_type_rows = [(q(slug(value)), q(value)) for value in mda_types]
    mda_rows = [
        (
            q(code),
            q(name),
            f"(select id from public.mda_types where name = {q(mda_type)})",
        )
        for code, name, mda_type in mdas
    ]
    return f"""insert into public.mda_types (slug, name)
values
{values_block(mda_type_rows)}
on conflict (slug) do update
set name = excluded.name,
    active = true,
    updated_at = now();

insert into public.mdas (code, name, mda_type_id)
values
{values_block(mda_rows)}
on conflict (code) do update
set name = excluded.name,
    mda_type_id = excluded.mda_type_id,
    active = true,
    updated_at = now();
"""


def lga_facility_seed(facility_ws) -> str:
    pairs = []
    for row in range(3, facility_ws.max_row + 1):
        lga = facility_ws.cell(row, 2).value
        facility = facility_ws.cell(row, 3).value
        if not lga or not facility:
            continue
        if str(lga).upper().startswith("GRAND TOTAL"):
            continue
        pairs.append((str(lga).strip(), str(facility).strip()))

    lgas = sorted({lga for lga, _ in pairs})
    lga_rows = [(q(name),) for name in lgas]
    facility_rows = [
        (
            f"(select id from public.lgas where name = {q(lga)})",
            q(facility),
            q("phc"),
        )
        for lga, facility in pairs
    ]

    return f"""insert into public.lgas (name)
values
{values_block(lga_rows)}
on conflict (name) do update
set active = true,
    updated_at = now();

insert into public.facilities (lga_id, name, facility_type)
values
{values_block(facility_rows)}
on conflict (lga_id, name) do update
set facility_type = excluded.facility_type,
    active = true,
    updated_at = now();
"""


def budget_seed(budget_ws) -> str:
    rows = []
    for row in range(5, budget_ws.max_row + 1):
        code = budget_ws.cell(row, 1).value
        name = budget_ws.cell(row, 2).value
        if not code or not name:
            continue
        if "SECTOR TOTAL" in str(name).upper():
            continue
        personnel = budget_ws.cell(row, 3).value
        other_recurrent = budget_ws.cell(row, 4).value
        total_recurrent = budget_ws.cell(row, 5).value
        capital = budget_ws.cell(row, 6).value
        total = budget_ws.cell(row, 7).value
        rows.append(
            (
                "2026",
                f"(select id from public.mdas where code = {q(str(code).strip())})",
                money(personnel),
                money(other_recurrent),
                money(total_recurrent),
                money(capital),
                money(total),
                q("Kano State Government 2026 Approved Budget"),
            )
        )

    return f"""insert into public.approved_budgets (
  fiscal_year,
  mda_id,
  personnel_amount,
  other_recurrent_amount,
  total_recurrent_amount,
  capital_amount,
  total_budget_amount,
  source_label
)
values
{values_block(rows)}
on conflict (fiscal_year, mda_id) do update
set personnel_amount = excluded.personnel_amount,
    other_recurrent_amount = excluded.other_recurrent_amount,
    total_recurrent_amount = excluded.total_recurrent_amount,
    capital_amount = excluded.capital_amount,
    total_budget_amount = excluded.total_budget_amount,
    source_label = excluded.source_label,
    updated_at = now();
"""


def aop_seed(aop_ws) -> str:
    rows = []
    for row in range(4, aop_ws.max_row + 1):
        code = aop_ws.cell(row, 1).value
        description = aop_ws.cell(row, 2).value
        mda = aop_ws.cell(row, 3).value
        cost = aop_ws.cell(row, 4).value
        if not code or not description or not mda or cost in (None, ""):
            continue
        rows.append(
            (
                "2026",
                q(str(code).strip()),
                q(str(description).strip()),
                f"(select id from public.mdas where name = {q(str(mda).strip())})",
                money(cost),
                str(row),
            )
        )

    return f"""insert into public.aop_activities (
  fiscal_year,
  activity_code,
  description,
  mda_id,
  budgeted_cost,
  source_row_number
)
values
{values_block(rows)}
on conflict (fiscal_year, activity_code, mda_id, source_row_number) do update
set description = excluded.description,
    budgeted_cost = excluded.budgeted_cost,
    source_row_number = excluded.source_row_number,
    active = true,
    updated_at = now();
"""


def main() -> None:
    wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
    lookup = wb["LOOKUP"]

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    reference_sql = "\n".join(
        [
            "-- Generated from Kano Health Finance Tracker.xlsx. Do not edit by hand.",
            "begin;",
            status_seed(),
            mda_seed(lookup),
            reference_insert("funding_sources", col_values(lookup, 4)),
            reference_insert("expenditure_categories", col_values(lookup, 6)),
            reference_insert("programme_areas", col_values(lookup, 8)),
            reference_insert("payment_methods", col_values(lookup, 18)),
            lga_facility_seed(wb["Facility Analysis"]),
            "commit;",
            "",
        ]
    )
    (OUTPUT_DIR / "001_reference_data.sql").write_text(reference_sql)

    budget_sql = "\n".join(
        [
            "-- Generated from Kano Health Finance Tracker.xlsx. Do not edit by hand.",
            "begin;",
            budget_seed(wb["Budget 2026"]),
            "commit;",
            "",
        ]
    )
    (OUTPUT_DIR / "002_budget_2026.sql").write_text(budget_sql)

    aop_sql = "\n".join(
        [
            "-- Generated from Kano Health Finance Tracker.xlsx. Do not edit by hand.",
            "begin;",
            aop_seed(wb["AOP 2026"]),
            "commit;",
            "",
        ]
    )
    (OUTPUT_DIR / "003_aop_2026.sql").write_text(aop_sql)

    print("Generated:")
    print(f"- {OUTPUT_DIR / '001_reference_data.sql'}")
    print(f"- {OUTPUT_DIR / '002_budget_2026.sql'}")
    print(f"- {OUTPUT_DIR / '003_aop_2026.sql'}")


if __name__ == "__main__":
    main()
