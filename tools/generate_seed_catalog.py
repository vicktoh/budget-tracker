from __future__ import annotations

from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl


WORKBOOK_PATH = Path("/Users/kunle/Downloads/Kano Health Finance Tracker.xlsx")
OUTPUT_PATH = Path("docs/seed-data-catalog.md")


def text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def money(value) -> str:
    if value is None or value == "":
        return "0.00"
    dec = Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return f"{dec:,.2f}"


def col_values(ws, col_idx: int) -> list[str]:
    values = []
    for row in range(2, ws.max_row + 1):
        value = text(ws.cell(row, col_idx).value)
        if value:
            values.append(value)
    return values


def infer_mda_type(name: str) -> str:
    lowered = name.lower()
    if "ministry" in lowered:
        return "Ministry"
    if "board" in lowered:
        return "Board"
    if "hospital" in lowered:
        return "Hospital"
    if "agency" in lowered:
        return "Agency"
    if "college" in lowered:
        return "College"
    if "school" in lowered:
        return "School"
    if "trust fund" in lowered or "fund" in lowered:
        return "Fund"
    if "centre" in lowered or "center" in lowered:
        return "Centre"
    return "Other"


def table(headers: list[str], rows: list[list[str]]) -> str:
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    for row in rows:
        escaped = [value.replace("|", "\\|") for value in row]
        lines.append("| " + " | ".join(escaped) + " |")
    return "\n".join(lines)


def main() -> None:
    wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
    lookup = wb["LOOKUP"]
    budget_ws = wb["Budget 2026"]
    facility_ws = wb["Facility Analysis"]

    mdas = []
    for row in range(2, lookup.max_row + 1):
        code = text(lookup.cell(row, 1).value)
        name = text(lookup.cell(row, 2).value)
        if code and name:
            mdas.append([code, name, infer_mda_type(name)])

    mda_types = sorted({row[2] for row in mdas})
    funding_sources = col_values(lookup, 4)
    expenditure_categories = col_values(lookup, 6)
    programme_areas = col_values(lookup, 8)
    payment_methods = col_values(lookup, 18)

    facility_pairs = []
    for row in range(3, facility_ws.max_row + 1):
        lga = text(facility_ws.cell(row, 2).value)
        facility = text(facility_ws.cell(row, 3).value)
        if not lga or not facility or lga.upper().startswith("GRAND TOTAL"):
            continue
        facility_pairs.append((lga, facility))

    lgas = sorted({lga for lga, _ in facility_pairs})
    facilities_by_lga = {lga: [] for lga in lgas}
    for lga, facility in facility_pairs:
        facilities_by_lga[lga].append(facility)

    budgets = []
    for row in range(5, budget_ws.max_row + 1):
        code = text(budget_ws.cell(row, 1).value)
        name = text(budget_ws.cell(row, 2).value)
        if not code or not name or "SECTOR TOTAL" in name.upper():
            continue
        budgets.append(
            [
                code,
                name,
                money(budget_ws.cell(row, 3).value),
                money(budget_ws.cell(row, 4).value),
                money(budget_ws.cell(row, 5).value),
                money(budget_ws.cell(row, 6).value),
                money(budget_ws.cell(row, 7).value),
            ]
        )

    lines = [
        "# Seed Data Catalog",
        "",
        "Source workbook: `/Users/kunle/Downloads/Kano Health Finance Tracker.xlsx`.",
        "",
        "This catalog enumerates the data seeded into Supabase before routine user entry. AOP activities are intentionally excluded because they are documented and seeded separately.",
        "",
        "## Summary",
        "",
        table(
            ["Seed area", "Target table", "Count", "Feeds dropdowns?"],
            [
                ["Entry statuses", "`entry_statuses`", "4", "Yes, reviewer/admin status filters and workflow controls"],
                ["MDA types", "`mda_types`", str(len(mda_types)), "Yes, admin/reference filters"],
                ["MDAs", "`mdas`", str(len(mdas)), "Yes, entry forms, dashboards, imports, filters"],
                ["Funding sources", "`funding_sources`", str(len(funding_sources)), "Yes, funding form and filters"],
                ["Expenditure categories", "`expenditure_categories`", str(len(expenditure_categories)), "Yes, expenditure form and filters"],
                ["Programme areas", "`programme_areas`", str(len(programme_areas)), "Yes, funding/expenditure forms and filters"],
                ["Payment methods", "`payment_methods`", str(len(payment_methods)), "Yes, expenditure form and filters"],
                ["LGAs", "`lgas`", str(len(lgas)), "Yes, PHC expenditure form and filters"],
                ["PHC facilities", "`facilities`", str(len(facility_pairs)), "Yes, dependent PHC facility dropdown"],
                ["2026 approved budgets", "`approved_budgets`", str(len(budgets)), "No direct entry dropdown; feeds dashboards and budget filters"],
            ],
        ),
        "",
        "## Dropdown Field Map",
        "",
        table(
            ["Form/report field", "Seeded table", "Seeded values used"],
            [
                ["MDA", "`mdas`", "21 MDA names, each with code and MDA type"],
                ["MDA Type", "`mda_types`", "8 inferred MDA classifications"],
                ["Funding Source", "`funding_sources`", "8 workbook values"],
                ["Expenditure Category", "`expenditure_categories`", "11 workbook values"],
                ["Programme Area", "`programme_areas`", "24 workbook values"],
                ["Payment Method", "`payment_methods`", "6 workbook values"],
                ["Entry Status", "`entry_statuses`", "4 workflow statuses"],
                ["LGA", "`lgas`", "44 Kano LGAs"],
                ["Facility Name", "`facilities`", "471 PHC facilities filtered by selected LGA"],
                ["Expenditure Item", "`expenditure_items`", "No initial seed from workbook; admin-managed after launch"],
                ["AOP Activity", "`aop_activities`", "Excluded from this catalog by request"],
            ],
        ),
        "",
        "## Entry Statuses",
        "",
        table(
            ["Slug", "Name", "Description"],
            [
                ["pending", "Pending", "Submitted by an MDA user and awaiting review"],
                ["approved", "Approved", "Reviewed and accepted for reporting"],
                ["processed", "Processed", "Accepted and reconciled or posted in the finance process"],
                ["rejected", "Rejected", "Reviewed and rejected; excluded from official reporting"],
            ],
        ),
        "",
        "## MDA Types",
        "",
        table(["Name"], [[value] for value in mda_types]),
        "",
        "## MDAs",
        "",
        table(["Code", "Name", "MDA Type"], mdas),
        "",
        "## Funding Sources",
        "",
        table(["Name"], [[value] for value in funding_sources]),
        "",
        "## Expenditure Categories",
        "",
        table(["Name"], [[value] for value in expenditure_categories]),
        "",
        "## Programme Areas",
        "",
        table(["Name"], [[value] for value in programme_areas]),
        "",
        "## Payment Methods",
        "",
        table(["Name"], [[value] for value in payment_methods]),
        "",
        "## LGAs",
        "",
        table(["Name"], [[value] for value in lgas]),
        "",
        "## PHC Facilities",
        "",
        "All seeded facilities use `facility_type = phc`. The facility dropdown should be filtered by the selected LGA.",
        "",
    ]

    for lga in lgas:
        lines.extend(
            [
                f"### {lga}",
                "",
                table(["Facility Name"], [[facility] for facility in facilities_by_lga[lga]]),
                "",
            ]
        )

    lines.extend(
        [
            "## 2026 Approved Budgets",
            "",
            "These rows feed budget-versus-actual dashboards and filters. The workbook sector-total row is not imported as an MDA budget row.",
            "",
            table(
                [
                    "MDA Code",
                    "MDA",
                    "Personnel",
                    "Other Recurrent",
                    "Total Recurrent",
                    "Capital",
                    "Total Budget",
                ],
                budgets,
            ),
            "",
            "## Not Seeded Initially",
            "",
            table(
                ["Data area", "Reason"],
                [
                    ["Expenditure items", "The workbook labels the field but does not provide an item list. Admins will manage this dropdown in the platform."],
                    ["Users and MDA memberships", "These depend on real Supabase Auth users and should be configured during onboarding."],
                    ["Submission windows", "The default v1 posture is permissive until admins create windows."],
                    ["Unspecified PHC Facility fallback rows", "The platform supports them, but they are not auto-seeded; admins can create them per LGA if needed."],
                    ["Entry attachments", "Created only when users upload supporting documents."],
                    ["Reference value requests", "Created by MDA users when a dropdown value is missing."],
                    ["Funding and expenditure entries", "Routine reporting data should be entered through authenticated forms or historical admin imports."],
                ],
            ),
            "",
        ]
    )

    OUTPUT_PATH.write_text("\n".join(lines))
    print(f"Wrote {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
