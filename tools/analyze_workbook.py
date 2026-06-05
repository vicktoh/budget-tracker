import json
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


WORKBOOK_PATH = Path("/Users/kunle/Downloads/Kano Health Finance Tracker.xlsx")


def cell_value(cell):
    value = cell.value
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def sheet_matrix(ws, max_rows=None, max_cols=None):
    rows = []
    last_row = min(ws.max_row, max_rows or ws.max_row)
    last_col = min(ws.max_column, max_cols or ws.max_column)
    for row in ws.iter_rows(min_row=1, max_row=last_row, min_col=1, max_col=last_col):
        values = [cell_value(cell) for cell in row]
        if any(value is not None for value in values):
            rows.append(values)
    return rows


def non_empty_rows(ws, max_rows=None):
    rows = []
    last_row = min(ws.max_row, max_rows or ws.max_row)
    for row in ws.iter_rows(min_row=1, max_row=last_row):
        values = [cell_value(cell) for cell in row]
        if any(value is not None for value in values):
            rows.append(
                {
                    "row": row[0].row,
                    "values": values,
                }
            )
    return rows


def formulas(ws):
    items = []
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                items.append({"cell": cell.coordinate, "formula": cell.value})
    return items


def data_validations(ws):
    validations = []
    for dv in ws.data_validations.dataValidation:
        validations.append(
            {
                "type": dv.type,
                "operator": dv.operator,
                "formula1": dv.formula1,
                "formula2": dv.formula2,
                "allow_blank": dv.allowBlank,
                "show_error_message": dv.showErrorMessage,
                "error_title": dv.errorTitle,
                "error": dv.error,
                "sqref": str(dv.sqref),
            }
        )
    return validations


def column_values(ws, col_idx):
    values = []
    for row in range(1, ws.max_row + 1):
        value = cell_value(ws.cell(row=row, column=col_idx))
        if value is not None:
            values.append(value)
    return values


def lookup_columns(ws):
    cols = {}
    for col_idx in range(1, ws.max_column + 1):
        values = column_values(ws, col_idx)
        if values:
            header = values[0]
            cols[get_column_letter(col_idx)] = {
                "header": header,
                "values": values[1:],
            }
    return cols


def defined_names(wb):
    names = {}
    for name, definition in wb.defined_names.items():
        destinations = []
        try:
            for title, coord in definition.destinations:
                destinations.append({"sheet": title, "range": coord})
        except Exception:
            destinations.append({"attr_text": definition.attr_text})
        names[name] = destinations
    return names


def tables(ws):
    return [
        {
            "name": table.name,
            "display_name": table.displayName,
            "ref": table.ref,
        }
        for table in ws.tables.values()
    ]


def header_row_guess(ws):
    best = None
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20)):
        values = [cell_value(cell) for cell in row]
        non_empty = [value for value in values if value is not None]
        if len(non_empty) >= 3 and (best is None or len(non_empty) > best["count"]):
            best = {"row": row[0].row, "count": len(non_empty), "values": values}
    return best


def main():
    wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=False)
    data_only = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)

    summary = {
        "workbook": str(WORKBOOK_PATH),
        "sheets": [],
        "defined_names": defined_names(wb),
    }

    for ws in wb.worksheets:
        ws_values = data_only[ws.title]
        sheet_info = {
            "title": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "sheet_state": ws.sheet_state,
            "tables": tables(ws),
            "header_guess": header_row_guess(ws),
            "data_validations": data_validations(ws),
            "formula_count": len(formulas(ws)),
            "formula_samples": formulas(ws)[:25],
        }
        if ws.title in {"Data Entry", "Funding Log", "Expenditure Log", "LOOKUP", "Budget 2026"}:
            sheet_info["non_empty_rows"] = non_empty_rows(ws, 80)
        if ws.title == "LOOKUP":
            sheet_info["lookup_columns"] = lookup_columns(ws)
        if ws.title in {"Dashboard", "Facility Analysis", "LGA Analysis"}:
            sheet_info["values_preview"] = sheet_matrix(ws_values, 40, 14)
        summary["sheets"].append(sheet_info)

    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
